"""知识库管理 API：上传文件、列出文档、删除文档。"""
from __future__ import annotations
import io

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile

from app.api.deps import require_permission
from app.api.schemas.models import KnowledgeIngestRequest
from app.rag.chunking import chunk_document
from app.workflows.nodes.rag_search import get_retriever

router = APIRouter(prefix="/knowledge", tags=["knowledge"])

_k_view = require_permission("knowledge.view")
_k_edit = require_permission("knowledge.edit")


# ---------- 工具：文件解析 ----------
def _parse_file(content_bytes: bytes, filename: str) -> str:
    """按扩展名解析文件为纯文本。"""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext == "pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content_bytes))
            return "\n\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"PDF 解析失败: {e}")

    if ext in ("docx", "doc"):
        try:
            from docx import Document
            doc = Document(io.BytesIO(content_bytes))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Word 解析失败: {e}")

    if ext in ("xlsx", "xls"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=True)
            lines = []
            for sheet in wb.worksheets:
                lines.append(f"# Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(c) if c is not None else "" for c in row]
                    if any(cells):
                        lines.append(" | ".join(cells))
            return "\n".join(lines)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Excel 解析失败: {e}")

    # txt / md / 其他 -> 纯文本
    try:
        return content_bytes.decode("utf-8")
    except UnicodeDecodeError:
        try:
            return content_bytes.decode("gbk")
        except Exception:
            return content_bytes.decode("utf-8", errors="replace")


# ---------- API ----------
@router.post("/ingest")
async def ingest(req: KnowledgeIngestRequest, user: dict = Depends(_k_edit)):
    """手动粘贴文本入库。"""
    chunks = chunk_document(req.content, doc_id=req.doc_id, source=req.source)
    retriever = get_retriever()
    retriever.index_chunks(chunks)
    # 知识库变更 -> 清空缓存
    try:
        from app.services.cache_service import get_cache
        cleared = get_cache().clear_all()
        print(f"[CACHE] cleared {cleared} entries after ingest")
    except Exception:
        pass
    return {
        "doc_id": req.doc_id,
        "chunks": len(chunks),
        "total_in_collection": retriever.count(),
    }


@router.post("/ingest-file")
async def ingest_file(
    file: UploadFile = File(...),
    doc_id: str = Form(...),
    source: str = Form(""),
    user: dict = Depends(_k_edit),
):
    """上传文件（PDF / Word / Excel / txt / md）并入库。"""
    if not file.filename:
        raise HTTPException(status_code=400, detail="文件名不能为空")

    content_bytes = await file.read()
    if not content_bytes:
        raise HTTPException(status_code=400, detail="文件内容为空")

    text = _parse_file(content_bytes, file.filename)
    if not text.strip():
        raise HTTPException(status_code=400, detail="文件解析后无有效文本")

    chunks = chunk_document(
        text,
        doc_id=doc_id,
        source=source or file.filename,
    )
    retriever = get_retriever()
    retriever.index_chunks(chunks)
    try:
        from app.services.cache_service import get_cache
        get_cache().clear_all()
    except Exception:
        pass

    return {
        "doc_id": doc_id,
        "filename": file.filename,
        "chars": len(text),
        "chunks": len(chunks),
        "total_in_collection": retriever.count(),
    }


@router.get("/docs")
async def list_docs(user: dict = Depends(_k_view)):
    """列出所有文档（按 doc_id 分组）。"""
    retriever = get_retriever()
    try:
        res = retriever.collection.get(include=["metadatas"])
    except Exception:
        return {"total": 0, "items": []}

    metas = res.get("metadatas", []) or []
    docs: dict = {}
    for m in metas:
        if not m:
            continue
        did = m.get("doc_id", "unknown")
        if did not in docs:
            docs[did] = {
                "doc_id": did,
                "source": m.get("source", ""),
                "parent_count": 0,
                "child_count": 0,
                "chunk_count": 0,
            }
        docs[did]["chunk_count"] += 1
        if m.get("type") == "parent":
            docs[did]["parent_count"] += 1
        else:
            docs[did]["child_count"] += 1

    items = sorted(docs.values(), key=lambda x: x["doc_id"])
    return {"total": len(items), "items": items}


@router.delete("/docs/{doc_id}")
async def delete_doc(doc_id: str, user: dict = Depends(_k_edit)):
    """删除某文档的所有切片。"""
    retriever = get_retriever()
    try:
        res = retriever.collection.get(include=["metadatas"])
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

    ids = res.get("ids", []) or []
    metas = res.get("metadatas", []) or []
    to_delete = [ids[i] for i, m in enumerate(metas)
                 if m and m.get("doc_id") == doc_id]

    if to_delete:
        retriever.collection.delete(ids=to_delete)

    return {"deleted": len(to_delete), "doc_id": doc_id}


# ============================================================
# 更新文档（编辑）
# ============================================================
from pydantic import BaseModel as _PM


class _UpdateDocReq(_PM):
    content: str
    source: str = ""
    new_doc_id: str = ""


@router.put("/docs/{doc_id}")
async def update_doc(doc_id: str, req: _UpdateDocReq,
                     user: dict = Depends(_k_edit)):
    """更新文档：删除旧的切片，用新内容重新入库。

    - content：新文本（必填）
    - source：来源（可选，不改则保留）
    - new_doc_id：想改 ID 时填（可选）
    """
    retriever = get_retriever()

    # 1. 读旧 metadatas 拿 source
    try:
        res = retriever.collection.get(include=["metadatas"])
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

    ids = res.get("ids", []) or []
    metas = res.get("metadatas", []) or []

    old_source = ""
    to_delete = []
    for i, m in enumerate(metas):
        if m and m.get("doc_id") == doc_id:
            to_delete.append(ids[i])
            if not old_source:
                old_source = m.get("source", "")

    if not to_delete:
        raise HTTPException(status_code=404, detail=f"文档 {doc_id} 不存在")

    final_source = req.source or old_source or "manual"
    final_doc_id = (req.new_doc_id or doc_id).strip()

    # 2. 删旧
    retriever.collection.delete(ids=to_delete)

    # 3. 新内容 chunk + index
    chunks = chunk_document(req.content, doc_id=final_doc_id, source=final_source)
    retriever.index_chunks(chunks)

    # 4. 清缓存
    try:
        from app.services.cache_service import get_cache
        get_cache().clear_all()
    except Exception:
        pass

    return {
        "old_doc_id": doc_id,
        "doc_id": final_doc_id,
        "source": final_source,
        "deleted": len(to_delete),
        "new_chunks": len(chunks),
        "total_in_collection": retriever.count(),
    }


@router.get("/stats")
async def stats(user: dict = Depends(_k_view)):
    retriever = get_retriever()
    return {"total_chunks": retriever.count()}

@router.get("/docs/{doc_id}/detail")
async def get_doc_detail(doc_id: str, user: dict = Depends(_k_view)):
    """返回某文档的完整内容（父块按顺序拼接）+ 所有切片。"""
    retriever = get_retriever()
    try:
        res = retriever.collection.get(include=["metadatas", "documents"])
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

    ids = res.get("ids", []) or []
    metas = res.get("metadatas", []) or []
    docs = res.get("documents", []) or []

    # 过滤出该 doc_id 的所有切片
    items = []
    for i, m in enumerate(metas):
        if not m or m.get("doc_id") != doc_id:
            continue
        items.append({
            "chunk_id": ids[i],
            "type": m.get("type"),
            "section_index": m.get("section_index", 0),
            "sub_index": m.get("sub_index", 0),
            "heading": m.get("heading_path", ""),
            "source": m.get("source", ""),
            "parent_id": m.get("parent_id", ""),
            "text": docs[i] or "",
        })

    if not items:
        raise HTTPException(status_code=404, detail=f"doc_id {doc_id} 不存在")

    # 父块按 section_index 排序拼接为完整内容
    parents = sorted(
        [x for x in items if x["type"] == "parent"],
        key=lambda x: x["section_index"],
    )
    full_text = "\n\n".join(p["text"] for p in parents)

    # 子块单独列出
    children = sorted(
        [x for x in items if x["type"] == "child"],
        key=lambda x: (x["section_index"], x["sub_index"]),
    )

    return {
        "doc_id": doc_id,
        "source": items[0].get("source", ""),
        "total_chunks": len(items),
        "parent_count": len(parents),
        "child_count": len(children),
        "full_text": full_text,
        "parents": parents,
        "children": children,
    }
